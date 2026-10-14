from __future__ import annotations

import fnmatch
import re
from pathlib import Path
from typing import Any

import pandas as pd
import pytest
from openpyxl import load_workbook  # type: ignore[import-untyped]

from core.config import load_config
from tests.fixtures.generate import LAB_A, LAB_B, LAB_C, LAB_D, generate
from tests.fixtures.truth import (
    ASSUMED_CONFIG,
    DEFAULT_SEED,
    REQUIRED_CASE_IDS,
    InjectedCase,
    Truth,
)

ROOT = Path(__file__).resolve().parents[1]
CONFIG_DIR = Path(__file__).parent / "fixtures" / "config"
ALIASES = {"MW-1", "MW1", "MW-01"}


def _xlsx_cells(path: Path) -> dict[str, list[tuple[Any, ...]]]:
    wb = load_workbook(path, data_only=True)
    return {ws.title: [tuple(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def _long_frame(project: Path) -> pd.DataFrame:
    a = pd.read_excel(project / LAB_A, dtype=str, keep_default_na=False)
    b = pd.read_csv(project / LAB_B, dtype=str, keep_default_na=False)
    return pd.concat([a, b], ignore_index=True).fillna("")


def _case(truth: Truth, case_id: str) -> InjectedCase:
    return next(c for c in truth.cases if c.case_id == case_id)


@pytest.fixture(scope="module")
def built(tmp_path_factory: pytest.TempPathFactory) -> tuple[Path, Truth]:
    out = tmp_path_factory.mktemp("gen")
    truth = generate("project_basic", DEFAULT_SEED, out)
    return out / "project_basic", truth


@pytest.mark.parametrize("case_id", REQUIRED_CASE_IDS)
def test_required_case_present(built: tuple[Path, Truth], case_id: str) -> None:
    project, truth = built
    on_disk = Truth.from_json((project / "truth.json").read_text(encoding="utf-8"))
    assert case_id in truth.case_ids()
    assert case_id in on_disk.case_ids()


def test_required_ids_unique() -> None:
    assert len(set(REQUIRED_CASE_IDS)) == len(REQUIRED_CASE_IDS)


def test_truth_json_roundtrip(built: tuple[Path, Truth]) -> None:
    project, truth = built
    text = (project / "truth.json").read_text(encoding="utf-8")
    assert Truth.from_json(text).to_json() == text
    assert truth.to_json() == text


def test_rerun_is_identical(tmp_path: Path) -> None:
    t1 = generate("project_basic", DEFAULT_SEED, tmp_path / "one")
    t2 = generate("project_basic", DEFAULT_SEED, tmp_path / "two")
    p1, p2 = tmp_path / "one" / "project_basic", tmp_path / "two" / "project_basic"
    assert t1 == t2
    names1 = sorted(p.name for p in p1.iterdir())
    assert names1 == sorted(p.name for p in p2.iterdir())
    assert (p1 / "truth.json").read_bytes() == (p2 / "truth.json").read_bytes()
    for name in names1:
        if name.endswith(".csv"):
            assert (p1 / name).read_bytes() == (p2 / name).read_bytes(), name
        elif name.endswith(".xlsx"):
            assert _xlsx_cells(p1 / name) == _xlsx_cells(p2 / name), name


def test_different_seed_changes_some_values(tmp_path: Path) -> None:
    t1 = generate("project_basic", 1, tmp_path / "s1")
    t2 = generate("project_basic", 2, tmp_path / "s2")
    assert t1.case_ids() == t2.case_ids()
    p1, p2 = tmp_path / "s1" / "project_basic", tmp_path / "s2" / "project_basic"
    changed = 0
    for name in (LAB_A, LAB_B, LAB_C, LAB_D):
        if name.endswith(".csv"):
            changed += (p1 / name).read_bytes() != (p2 / name).read_bytes()
        else:
            changed += _xlsx_cells(p1 / name) != _xlsx_cells(p2 / name)
    assert changed >= 1


def test_expected_counts_nonzero(built: tuple[Path, Truth]) -> None:
    _, truth = built
    for case in truth.cases:
        assert case.expected, case.case_id
        for key, value in case.expected.items():
            if isinstance(value, bool):
                continue
            if isinstance(value, (int, float)):
                assert value > 0, (case.case_id, key)
            elif isinstance(value, (list, tuple, dict, str)):
                assert len(value) > 0, (case.case_id, key)


def test_case_files_exist(built: tuple[Path, Truth]) -> None:
    project, truth = built
    for case in truth.cases:
        assert case.description, case.case_id
        assert case.files, case.case_id
        for f in case.files:
            assert (project / f).is_file(), (case.case_id, f)


def test_addendum_cases_tagged(built: tuple[Path, Truth]) -> None:
    _, truth = built
    tags = {
        "a12_nd_lt_zero": "A12",
        "a12_nd_dl_zero": "A12",
        "a5_gt_fails_min_obs": "A5",
        "a11_nd_differing_dls": "A11",
        "q9_qc_field_duplicate": "Q9",
    }
    for case_id, tag in tags.items():
        assert tag in _case(truth, case_id).description


def test_generated_dir_gitignored() -> None:
    lines = (ROOT / ".gitignore").read_text(encoding="utf-8").splitlines()
    assert "tests/fixtures/generated/" in {ln.strip() for ln in lines}


def test_writes_only_under_out_root(tmp_path: Path) -> None:
    out = tmp_path / "out"
    generate("project_basic", DEFAULT_SEED, out)
    assert [p.name for p in tmp_path.iterdir()] == ["out"]
    assert [p.name for p in out.iterdir()] == ["project_basic"]


def test_location_alias_rows_recount(built: tuple[Path, Truth]) -> None:
    project, truth = built
    df = _long_frame(project)
    case = _case(truth, "location_aliases")
    assert int(df["location"].isin(ALIASES).sum()) == case.expected["rows"]
    assert set(case.expected["locations"]) == ALIASES
    assert set(df.loc[df["location"].isin(ALIASES), "location"]) == ALIASES


def test_right_censored_recount(built: tuple[Path, Truth]) -> None:
    project, truth = built
    df = _long_frame(project)
    n = int(df["result"].str.strip().str.startswith(">").sum())
    expected = (
        _case(truth, "right_censored").expected["rows"]
        + _case(truth, "a5_gt_fails_min_obs").expected["gt_rows"]
    )
    assert n == expected


def test_lt_zero_recount(built: tuple[Path, Truth]) -> None:
    project, truth = built
    df = _long_frame(project)
    n = int((df["result"].str.strip() == "<0").sum())
    assert n == _case(truth, "a12_nd_lt_zero").expected["rows"]


def test_blank_result_and_dl_zero_recount(built: tuple[Path, Truth]) -> None:
    project, truth = built
    df = _long_frame(project)
    blank = df["result"].str.strip() == ""
    nd = df["detect_flag"].str.strip().str.upper() == ASSUMED_CONFIG.nd_detect_flag
    mdl = pd.to_numeric(df["mdl"].str.replace(",", ".", regex=False), errors="coerce")
    assert int((blank & nd & (mdl > 0)).sum()) == _case(truth, "blank_result_nd").expected["rows"]
    assert int((blank & nd & (mdl == 0)).sum()) == _case(truth, "a12_nd_dl_zero").expected["rows"]


def test_no_result_tokens_recount(built: tuple[Path, Truth]) -> None:
    project, truth = built
    df = _long_frame(project)
    tokens = {t.casefold() for t in ASSUMED_CONFIG.no_result_tokens}
    is_token = df["result"].str.strip().str.casefold().isin(tokens)
    no_dl = (df["mdl"].str.strip() == "") & (df["rl"].str.strip() == "")
    case = _case(truth, "no_result_tokens")
    assert int((is_token & no_dl).sum()) == case.expected["rows"]
    assert set(case.expected["tokens"]) == set(ASSUMED_CONFIG.no_result_tokens)


def test_lab_b_file_name_matches_date_rule_pattern() -> None:
    assert fnmatch.fnmatch(LAB_B, ASSUMED_CONFIG.lab_b_source_pattern)
    assert LAB_B == ASSUMED_CONFIG.lab_b_file_name


def test_a5_case_shape(built: tuple[Path, Truth]) -> None:
    project, truth = built
    case = _case(truth, "a5_gt_fails_min_obs")
    df = _long_frame(project)
    sel = df[
        (df["location"] == case.expected["location"]) & (df["analyte"] == case.expected["analyte"])
    ]
    res = sel["result"].str.strip()
    n_gt = int(res.str.startswith(">").sum())
    assert n_gt >= 1
    assert len(sel) >= ASSUMED_CONFIG.min_obs_per_cell
    assert len(sel) - n_gt < ASSUMED_CONFIG.min_obs_per_cell


def test_a11_case_shape(built: tuple[Path, Truth]) -> None:
    _, truth = built
    case = _case(truth, "a11_nd_differing_dls")
    dls = case.expected["dls"]
    assert len(case.expected["locations"]) >= 2
    assert len(set(dls)) >= 2
    assert min(dls) < 0.1 * max(dls)


def test_q9_field_duplicate_has_parent(built: tuple[Path, Truth]) -> None:
    project, truth = built
    df = _long_frame(project)
    case = _case(truth, "q9_qc_field_duplicate")
    dup = df[df["sample_id"] == case.expected["sample_id"]]
    assert len(dup) > 0
    assert (dup["parent_sample_id"] == case.expected["parent_sample_id"]).all()
    assert (df["sample_id"] == case.expected["parent_sample_id"]).any()


def test_assumed_config_values_in_fixture_config() -> None:
    cfg = load_config([CONFIG_DIR])
    a = ASSUMED_CONFIG

    codes = {k.upper() for k in cfg.flag_map}
    assert set(a.qualifier_codes) <= codes

    rules = [
        r
        for r in cfg.date_rules
        if r.source_pattern == a.lab_b_source_pattern and r.format == a.lab_b_date_format
    ]
    assert rules, "lab_b date rule missing"

    matched = [
        m for r in cfg.header_rules if (m := re.match(r.regex, a.header_example)) is not None
    ]
    assert matched, "no header rule matches the example header"
    m = matched[0]
    assert (m.group("analyte"), m.group("fraction"), m.group("unit")) == a.header_example_parts

    assert set(a.units) <= {u.unit for u in cfg.units}

    cfg_tokens = {t.casefold() for t in cfg.tokens.no_result_tokens}
    assert {t.casefold() for t in a.no_result_tokens} <= cfg_tokens

    assert set(a.analyte_keys) <= {r.key for r in cfg.analytes}
    assert cfg.qa_thresholds.min_obs_per_cell == a.min_obs_per_cell

    no3_n = [
        c
        for c in cfg.conversions
        if c.analyte == "no3" and c.from_unit == "mg/L as N" and c.to_unit == "mg/L"
    ]
    assert no3_n
    assert no3_n[0].factor == a.no3_as_n_factor

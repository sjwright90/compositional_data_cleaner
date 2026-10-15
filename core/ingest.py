"""Source inventory and raw sheet reading (spec §6.1; addendum A15, A17). Pure: no writes.

Supported formats are `.xlsx` and `.csv` only (A17). A csv has one sheet, named `""`.

`read_sheet` returns every cell as a string so parsing stays with T11/T12; the frame has integer
column labels (no header inference) and no trailing empty rows or columns. Cell rule:

- numeric cells -> shortest round-trip repr (`1.374` -> `"1.374"`, no float noise);
- date/datetime cells -> ISO 8601 date, with the time only when it is not midnight;
- booleans -> `"TRUE"` / `"FALSE"`;
- text unchanged (no trimming); empty cells -> `""`.

A csv is read as text exactly as written (UTF-8, with a BOM allowed; Windows-1252 if that fails).
"""

from __future__ import annotations

import csv
import logging
from collections.abc import Iterable, Iterator
from dataclasses import dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook  # type: ignore[import-untyped]

from core.errors import ProjectError
from core.pipeline import FileLockedError

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES: tuple[str, ...] = (".xlsx", ".csv")

# Rows shown in an inventory preview. A module constant, not config: it only sizes the UI preview.
PREVIEW_ROWS = 20


class UnsupportedSourceError(ProjectError):
    """A source file has a format the app does not read (A17)."""

    def __init__(self, paths: Iterable[Path | str]) -> None:
        self.paths = [str(p) for p in paths]
        listed = ", ".join(f"'{p}'" for p in self.paths)
        super().__init__(
            f"These files are not supported: {listed}. Only .xlsx and .csv files can be used; "
            "save it as .xlsx or .csv and try again. Nothing was imported."
        )


class SourceReadError(ProjectError):
    """A source file or sheet could not be read."""


@dataclass(frozen=True)
class SheetInfo:
    file: Path
    sheet: str  # "" for a csv
    rows: int
    cols: int
    preview: pd.DataFrame = field(compare=False, hash=False, repr=False)


def _check_suffix(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedSourceError([path])
    return suffix


def _locked(path: Path) -> FileLockedError:
    return FileLockedError(
        f"Cannot read '{path.name}'. If it is open in Excel, close it and try again."
    )


def _cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        return repr(value)
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat()
    if isinstance(value, date | time):
        return value.isoformat()
    return str(value)


def _trim(rows: list[list[str]]) -> list[list[str]]:
    """Drop trailing all-empty rows and columns; pad the rest to a rectangle."""
    while rows and not any(rows[-1]):
        rows.pop()
    width = max(
        (max((i for i, v in enumerate(r) if v != ""), default=-1) + 1 for r in rows), default=0
    )
    return [r[:width] + [""] * (width - len(r)) for r in rows]


def _xlsx_rows(path: Path) -> dict[str, list[list[str]]]:
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
    except PermissionError as exc:
        raise _locked(path) from exc
    except Exception as exc:  # unreadable user file: any parser failure is reported, not raised
        raise SourceReadError(
            f"'{path.name}' could not be read as an Excel workbook. The file may be damaged, or "
            "it may not really be an .xlsx file (for example a renamed .xls or .csv). Open it in "
            "Excel, save it as .xlsx and try again."
        ) from exc
    try:
        return {
            ws.title: _trim([[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)])
            for ws in wb.worksheets
        }
    except PermissionError as exc:
        raise _locked(path) from exc
    finally:
        wb.close()


def _csv_rows(path: Path) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            with path.open("r", encoding=encoding, newline="") as fh:
                return _trim([list(r) for r in csv.reader(fh)])
        except PermissionError as exc:
            raise _locked(path) from exc
        except UnicodeDecodeError:
            continue
        except csv.Error as exc:
            raise SourceReadError(
                f"'{path.name}' could not be read as a csv file ({exc}). The file may be damaged."
            ) from exc
    raise SourceReadError(
        f"'{path.name}' could not be read: its text encoding is not recognised. "
        "Open it in Excel and save it as .xlsx or .csv (UTF-8)."
    )


def _frame(rows: list[list[str]]) -> pd.DataFrame:
    return pd.DataFrame(rows, dtype=str) if rows else pd.DataFrame(dtype=str)


def _sheets(path: Path) -> Iterator[tuple[str, list[list[str]]]]:
    if _check_suffix(path) == ".csv":
        yield "", _csv_rows(path)
    else:
        yield from _xlsx_rows(path).items()


def read_inventory(path: Path) -> tuple[SheetInfo, ...]:
    """One `SheetInfo` per sheet: name, dimensions and the first `PREVIEW_ROWS` rows."""
    path = Path(path)
    out = tuple(
        SheetInfo(path, sheet, len(rows), len(rows[0]) if rows else 0, _frame(rows[:PREVIEW_ROWS]))
        for sheet, rows in _sheets(path)
    )
    logger.info("inventory %s: %d sheet(s)", path.name, len(out))
    return out


def read_sheet(path: Path, sheet: str) -> pd.DataFrame:
    """Every cell of one sheet as a string (see the module docstring for the cell rule)."""
    path = Path(path)
    available: list[str] = []
    for name, rows in _sheets(path):
        if name == sheet:
            return _frame(rows)
        available.append(name)
    listed = ", ".join(f"'{n}'" for n in available) or "none"
    raise SourceReadError(
        f"'{path.name}' has no sheet named '{sheet}'. Sheets in this file: {listed}."
    )
